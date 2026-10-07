"""Uji ujung-ke-ujung: bangun deck sungguhan dari 12 berkas md + Design Style dengan snapshot fixture, lalu periksa isinya."""
import copy
import datetime as dt
import json
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
BUILD = os.path.dirname(HERE)
FOLDER = os.path.dirname(BUILD)
sys.path.insert(0, BUILD)

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

from deckgen import build, draw, theme  # noqa: E402

FIXTURE = os.path.join(HERE, "fixtures", "snapshot_min.json")
EMU = 914400


def _read(path):
    with open(path, encoding="utf-8") as f:
        return f.read()


def _write(path, text):
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)


class BuildE2E(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="deck-e2e-")
        cls.out = os.path.join(cls.tmp, "deck.pptx")
        cls.rep = build.build(FOLDER, cls.out, snapshot=FIXTURE)
        cls.prs = Presentation(cls.out)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_twelve_slides_one_per_markdown_file(self):
        docs, style = build.collect(FOLDER)
        self.assertEqual(len(docs), 12)
        self.assertIsNotNone(style, "99 - Design Style.md harus ada")
        self.assertEqual(len(self.prs.slides), 12)
        self.assertEqual(self.rep.n_slides, 12)

    def test_no_overflow_or_shrink_warnings(self):
        self.assertEqual(self.rep.warnings, [])

    def test_canvas_is_16_9(self):
        self.assertAlmostEqual(self.prs.slide_width / EMU, 13.333, places=2)
        self.assertAlmostEqual(self.prs.slide_height / EMU, 7.5, places=2)

    def test_every_slide_has_a_real_title_and_speaker_notes(self):
        for i, s in enumerate(self.prs.slides, 1):
            self.assertIsNotNone(s.shapes.title, f"slide {i} tanpa placeholder judul")
            self.assertTrue(s.shapes.title.text_frame.text.strip(), f"slide {i} judul kosong")
            notes = s.notes_slide.notes_text_frame.text
            for part in ("VISUAL SCRIPT", "NARRATION", "NUMBERS AS OF"):
                self.assertIn(part, notes, f"slide {i}: catatan tanpa {part}")

    def test_text_boxes_stay_inside_the_slide(self):
        w, h = self.prs.slide_width, self.prs.slide_height
        tol = int(0.02 * EMU)
        for i, s in enumerate(self.prs.slides, 1):
            for sh in s.shapes:
                if getattr(sh, "has_text_frame", False) and sh.has_text_frame and sh.text_frame.text.strip():
                    self.assertGreaterEqual(sh.left, -tol, f"slide {i} '{sh.name}' keluar kiri")
                    self.assertGreaterEqual(sh.top, -tol, f"slide {i} '{sh.name}' keluar atas")
                    self.assertLessEqual(sh.left + sh.width, w + tol, f"slide {i} '{sh.name}' keluar kanan")
                    self.assertLessEqual(sh.top + sh.height, h + tol, f"slide {i} '{sh.name}' keluar bawah")

    def test_only_token_fonts_are_used(self):
        fonts = set(theme.DEFAULT_TOKENS["fonts"]["brand"].values())
        used = set()
        for s in self.prs.slides:
            for el in s._element.iter(qn("a:latin")):
                used.add(el.get("typeface"))
        self.assertTrue(used <= fonts, f"font di luar token: {used - fonts}")
        self.assertGreaterEqual(len(used), 3)

    def test_all_runs_are_tagged_english(self):
        langs = {r.get("lang") for s in self.prs.slides for r in s._element.iter(qn("a:rPr")) if r.get("lang")}
        self.assertEqual(langs, {"en-US"})

    def test_pictures_have_alt_text(self):
        pics = [(i, sh) for i, s in enumerate(self.prs.slides, 1) for sh in s.shapes if sh.shape_type == 13]
        self.assertGreaterEqual(len(pics), 2)
        for i, sh in pics:
            self.assertTrue(sh._element.xpath(".//p:cNvPr/@descr"), f"slide {i}: gambar tanpa teks alternatif")

    def test_numbers_come_from_the_snapshot_not_from_typing(self):
        text9 = " ".join(sh.text_frame.text for sh in self.prs.slides[8].shapes if getattr(sh, "has_text_frame", False) and sh.has_text_frame)
        f = self.rep.facts
        for needle in (str(f["verified"]), str(f["silent"]), str(f["locks"]), f["block"]):
            self.assertIn(needle, text9)
        self.assertNotIn("{{", text9)

    def test_slide_numbers_use_a_field_and_notes_include_addresses(self):
        for s in self.prs.slides:
            self.assertTrue(list(s._element.iter(qn("a:fld"))), "nomor slide harus berupa field")
        notes10 = self.prs.slides[9].notes_slide.notes_text_frame.text
        self.assertIn("0x9B78200beFbbBe836585d31bd5b6dB32587064f3", notes10)

    def test_slides_use_fade_transition(self):
        for s in self.prs.slides:
            self.assertIsNotNone(s._element.find(qn("p:transition")))

    def test_core_properties_are_ours_not_the_template_default(self):
        cp = self.prs.core_properties
        self.assertEqual(cp.title, "Fabius: Submission Deck")
        self.assertEqual(cp.author, "Fabius")
        self.assertNotIn("Canny", cp.author + cp.last_modified_by)

    def test_theme_carries_fabius_palette_and_fonts(self):
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        blob = self.prs.slide_master.part.part_related_by(RT.THEME).blob.decode("utf-8")
        for needle in ("0F0B26", "6E4BFF", "Archivo", "Inter"):
            self.assertIn(needle, blob)


class BuildVariants(unittest.TestCase):
    def _copy_folder(self):
        td = tempfile.mkdtemp(prefix="deck-var-")
        dst = os.path.join(td, "13")
        shutil.copytree(FOLDER, dst, ignore=shutil.ignore_patterns("build", "__pycache__"))
        return td, dst

    def test_safe_fonts_profile(self):
        with tempfile.TemporaryDirectory() as td:
            out = os.path.join(td, "safe.pptx")
            build.build(FOLDER, out, snapshot=FIXTURE, font_profile="safe", only={1, 5})
            prs = Presentation(out)
            used = {el.get("typeface") for s in prs.slides for el in s._element.iter(qn("a:latin"))}
            self.assertEqual(used, {"Arial", "Courier New"})
            self.assertEqual(len(prs.slides), 2)

    def test_guard_failure_stops_the_build(self):
        snap = json.loads(_read(FIXTURE))
        snap["confidence"]["B1-TREND"]["fd16"] = "LOLOS"
        with tempfile.TemporaryDirectory() as td:
            p = os.path.join(td, "snap.json")
            _write(p, json.dumps(snap))
            with self.assertRaises(build.BuildError) as cm:
                build.build(FOLDER, os.path.join(td, "x.pptx"), snapshot=p, check_only=True)
            self.assertIn("11 - Built to Be Doubted.md", str(cm.exception))
            self.assertIn("guard gagal", str(cm.exception))

    def test_unknown_layout_is_reported(self):
        td, dst = self._copy_folder()
        try:
            p = os.path.join(dst, "02 - The Problem.md")
            s = _read(p).replace("layout: claims_gap", "layout: nonexistent")
            _write(p, s)
            with self.assertRaises(build.BuildError) as cm:
                build.build(dst, os.path.join(td, "x.pptx"), snapshot=FIXTURE, check_only=True)
            self.assertIn("nonexistent", str(cm.exception))
        finally:
            shutil.rmtree(td, ignore_errors=True)

    def test_missing_slide_number_is_reported(self):
        td, dst = self._copy_folder()
        try:
            os.remove(os.path.join(dst, "07 - The AI Desk.md"))
            with self.assertRaises(build.BuildError) as cm:
                build.build(dst, os.path.join(td, "x.pptx"), snapshot=FIXTURE, check_only=True)
            self.assertIn("berurutan", str(cm.exception))
        finally:
            shutil.rmtree(td, ignore_errors=True)

    def test_design_style_tokens_drive_the_build(self):
        """Mengubah token di 99 - Design Style.md harus mengubah deck (bukan hanya dokumentasi)."""
        td, dst = self._copy_folder()
        try:
            p = os.path.join(dst, "99 - Design Style.md")
            s = _read(p)
            self.assertIn('night: "0F0B26"', s)
            s = s.replace('night: "0F0B26"', 'night: "123456"').replace('footer: "FABIUS · SUBMISSION DECK"', 'footer: "TEST FOOTER"')
            _write(p, s)
            out = os.path.join(td, "x.pptx")
            build.build(dst, out, snapshot=FIXTURE, only={1})
            prs = Presentation(out)
            from pptx.opc.constants import RELATIONSHIP_TYPE as RT
            blob = prs.slide_master.part.part_related_by(RT.THEME).blob.decode("utf-8")
            self.assertIn("123456", blob)
            texts = " ".join(sh.text_frame.text for sh in prs.slides[0].shapes if getattr(sh, "has_text_frame", False) and sh.has_text_frame)
            self.assertIn("TEST FOOTER", texts.upper())
        finally:
            shutil.rmtree(td, ignore_errors=True)

    def test_unknown_fact_token_is_reported(self):
        td, dst = self._copy_folder()
        try:
            p = os.path.join(dst, "09 - Live Product.md")
            s = _read(p).replace('value: "{{verified}}"', 'value: "{{does_not_exist}}"', 1)
            _write(p, s)
            with self.assertRaises(build.BuildError) as cm:
                build.build(dst, os.path.join(td, "x.pptx"), snapshot=FIXTURE, check_only=True)
            self.assertIn("does_not_exist", str(cm.exception))
        finally:
            shutil.rmtree(td, ignore_errors=True)

    def test_missing_screenshot_is_reported(self):
        td, dst = self._copy_folder()
        try:
            os.remove(os.path.join(dst, "assets", "ui-home-desktop.jpg"))
            with self.assertRaises(build.BuildError) as cm:
                build.build(dst, os.path.join(td, "x.pptx"), snapshot=FIXTURE)
            self.assertIn("gambar tidak ditemukan", str(cm.exception))
        finally:
            shutil.rmtree(td, ignore_errors=True)

    def test_non_integer_slide_number_is_reported(self):
        td, dst = self._copy_folder()
        try:
            p = os.path.join(dst, "02 - The Problem.md")
            s = _read(p)
            self.assertIn("slide: 2\n", s)
            _write(p, s.replace("slide: 2\n", "slide: two\n", 1))
            with self.assertRaises(build.BuildError) as cm:
                build.build(dst, os.path.join(td, "x.pptx"), snapshot=FIXTURE, check_only=True)
            self.assertIn("bukan bilangan bulat", str(cm.exception))
        finally:
            shutil.rmtree(td, ignore_errors=True)

    def test_file_timestamp_is_not_frozen(self):
        now = dt.datetime.now(dt.timezone.utc).replace(tzinfo=None)
        self.assertLess(abs((now - draw.build_time()).total_seconds()), 120)
        old = os.environ.get("SOURCE_DATE_EPOCH")
        os.environ["SOURCE_DATE_EPOCH"] = "1790000000"
        try:
            self.assertEqual(draw.build_time(), dt.datetime(2026, 9, 21, 14, 13, 20))
        finally:
            if old is None:
                del os.environ["SOURCE_DATE_EPOCH"]
            else:
                os.environ["SOURCE_DATE_EPOCH"] = old
        self.assertEqual(draw.build_time(dt.datetime(2030, 1, 2, 3, 4, 5)), dt.datetime(2030, 1, 2, 3, 4, 5))

    def test_stale_render_output_is_removed(self):
        with tempfile.TemporaryDirectory() as td:
            for name in ("slide-01.png", "slide-13.png", "deck.pdf", "logo.png", "notes.txt"):
                _write(os.path.join(td, name), "x")
            build._hapus_basi(td, os.path.join(td, "deck.pdf"))
            self.assertEqual(sorted(os.listdir(td)), ["logo.png", "notes.txt"])

    @unittest.skipUnless(os.environ.get("DECK_RENDER") == "1", "set DECK_RENDER=1 (butuh soffice + pdftoppm, ±40 dtk)")
    def test_render_to_png(self):
        with tempfile.TemporaryDirectory() as td:
            out = os.path.join(td, "deck.pptx")
            build.build(FOLDER, out, snapshot=FIXTURE)
            imgs = build.render_images(out, os.path.join(td, "png"))
            self.assertEqual(len(imgs), 12)


if __name__ == "__main__":
    unittest.main()
