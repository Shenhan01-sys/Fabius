"""Tes unit: penanda teks, pengukur teks, pembaca md, fakta + guard, urutan XML DrawingML."""
import json
import os
import sys
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

from lxml import etree  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402
from pptx.util import Inches  # noqa: E402

from deckgen import facts, markup, mdparse, measure, xmlfx  # noqa: E402

FIXTURE = os.path.join(HERE, "fixtures", "snapshot_min.json")


class MarkupTests(unittest.TestCase):
    def test_accent_bold_mono_and_plain(self):
        runs = markup.parse_line("a **b** *c* `d` e")
        self.assertEqual([(r.text, r.bold, r.accent, r.mono) for r in runs],
                         [("a ", False, False, False), ("b", True, False, False), (" ", False, False, False), ("c", False, True, False),
                          (" ", False, False, False), ("d", False, False, True), (" e", False, False, False)])

    def test_newline_makes_lines(self):
        self.assertEqual(len(markup.parse("one\n*two*")), 2)

    def test_plain_strips_markers(self):
        self.assertEqual(markup.plain("**a** *b* `c`"), "a b c")

    def test_lone_asterisk_is_literal(self):
        self.assertEqual(markup.plain("5 * 3"), "5 * 3")


class MeasureTests(unittest.TestCase):
    def test_wider_text_needs_more_lines(self):
        t = "word " * 40
        self.assertGreater(measure.wrap_line_count(t, 2.0, 14), measure.wrap_line_count(t, 6.0, 14))

    def test_fit_size_shrinks_until_it_fits(self):
        s, ok = measure.fit_size("a long sentence " * 6, 3.0, 0.6, 20, 9)
        self.assertTrue(ok)
        self.assertLess(s, 20)

    def test_fit_size_reports_impossible(self):
        _, ok = measure.fit_size("x " * 400, 1.0, 0.2, 20, 10)
        self.assertFalse(ok)

    def test_balance_splits_two_line_text_evenly(self):
        t = "Plug in your bot or agent, prove every call on-chain, sell it to humans and AI via x402 + MCP."
        out = measure.balance(t, 6.6, 12.5)
        self.assertIn("\n", out)
        a, b = out.split("\n")
        self.assertLess(abs(len(a) - len(b)), 20)

    def test_balance_leaves_single_line_alone(self):
        self.assertEqual(measure.balance("short", 6.0, 12), "short")


class MdParseTests(unittest.TestCase):
    TEXT = "---\ntags: [a]\nslide: 2\n---\n# T\n\n## Scene\nhello\n\n## Slide\n```yaml\nk: v\nn: 3\n```\n\n## Narration\nsay it\n"

    def test_front_sections_and_yaml(self):
        d = mdparse.parse_text(self.TEXT)
        self.assertEqual(d.front["slide"], 2)
        self.assertEqual(d.data["slide"], {"k": "v", "n": 3})
        self.assertEqual(d.sections["narration"].strip(), "say it")
        self.assertEqual(d.title, "T")

    def test_missing_front_matter_is_an_error(self):
        with self.assertRaises(mdparse.MdError):
            mdparse.parse_text("# no front matter\n")

    def test_bad_yaml_is_an_error(self):
        with self.assertRaises(mdparse.MdError):
            mdparse.parse_text("---\nslide: 1\n---\n## Slide\n```yaml\nk: [unclosed\n```\n")

    def test_plain_notes_strip_markdown_and_wikilinks(self):
        self.assertEqual(mdparse.md_to_plain("**b** [[x/y z|alias]] `c`"), "b alias c")


class FactsTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.f = facts.compute(facts.load(FIXTURE))

    def test_core_numbers_are_consistent(self):
        f = self.f
        self.assertEqual(f["commits"], f["verified"] + f["alarms"] + f["missed"] + f["unopened"] + 0 * f["before_lock"])
        self.assertEqual(f["silent"] + f["signal_days"], f["verified"])
        self.assertEqual(f["bots_not_passed"], f["bots"] - f["fd16_passed"])

    def test_render_replaces_tokens(self):
        self.assertEqual(facts.render("{{commits}} / {{bots}}", self.f), f"{self.f['commits']} / {self.f['bots']}")

    def test_unknown_token_is_an_error(self):
        with self.assertRaises(facts.FactError):
            facts.render("{{nope}}", self.f)

    def test_guard_true_and_false(self):
        ok, _ = facts.check_guard("fd16_passed == 0", self.f)
        self.assertTrue(ok)
        bad = dict(self.f, fd16_passed=1)
        ok, msg = facts.check_guard("fd16_passed == 0", bad)
        self.assertFalse(ok)
        self.assertIn("sekarang 1", msg)

    def test_guard_rejects_unsafe_expressions(self):
        for g in ("__import__('os')", "commits == commits", "commits >= 1 and alarms == 0", "nope == 1"):
            with self.assertRaises(facts.FactError):
                facts.check_guard(g, self.f)

    def test_guard_only_applies_to_integer_facts(self):
        for g in ("lag_min_h >= 1", "asof == 3"):
            with self.assertRaises(facts.FactError):
                facts.check_guard(g, self.f)

    def test_short_address(self):
        self.assertEqual(facts.short_addr("0x9B78200beFbbBe836585d31bd5b6dB32587064f3"), "0x9B78…64f3")


class XmlFxTests(unittest.TestCase):
    def _shape(self):
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        fb = s.shapes.build_freeform(0, 0, scale=1.0)
        fb.add_line_segments([(Inches(1), 0), (Inches(1), Inches(1))], close=True)
        return fb.convert_to_shape()

    def test_child_order_follows_the_schema(self):
        sh = self._shape()
        xmlfx.effects(sh, glow=("9D86FF", 8, 0.4), shadow=("000000", 10, 4, 0.3))   # efek lebih dulu...
        xmlfx.line(sh, "FFFFFF", 1.0, 0.5)                                         # ...lalu garis...
        xmlfx.solid(sh, "6E4BFF", 0.5)                                              # ...lalu isi: urutan akhir harus tetap benar
        tags = [etree.QName(c).localname for c in sh._element.spPr]
        order = ["xfrm", "custGeom", "solidFill", "ln", "effectLst"]
        self.assertEqual(tags, order)

    def test_glow_comes_before_shadow(self):
        sh = self._shape()
        xmlfx.effects(sh, glow=("9D86FF", 8, 0.4), shadow=("000000", 10, 4, 0.3))
        eff = sh._element.spPr.find(qn("a:effectLst"))
        self.assertEqual([etree.QName(c).localname for c in eff], ["glow", "outerShdw"])

    def test_alpha_is_written_only_when_needed(self):
        sh = self._shape()
        xmlfx.solid(sh, "FFFFFF", 1.0)
        self.assertIsNone(sh._element.spPr.find(".//" + qn("a:alpha")))
        xmlfx.solid(sh, "FFFFFF", 0.25)
        self.assertEqual(sh._element.spPr.find(".//" + qn("a:alpha")).get("val"), "25000")

    def test_fill_replaces_fill(self):
        sh = self._shape()
        xmlfx.solid(sh, "FFFFFF")
        xmlfx.gradient(sh, [(0, "FFFFFF", 0.5), (100, "000000", 0.0)])
        sp = sh._element.spPr
        self.assertEqual(len([c for c in sp if etree.QName(c).localname in ("solidFill", "gradFill")]), 1)


if __name__ == "__main__":
    unittest.main()
