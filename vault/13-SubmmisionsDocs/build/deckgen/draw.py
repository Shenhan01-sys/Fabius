"""Primitif gambar deck: presentasi dasar (16:9 + tema Fabius), latar, teks, kartu kaca, pil status, penanda bernomor, gambar, catatan.

Semua bentuk adalah objek asli PowerPoint (bisa dipilih dan diedit), bukan gambar mentah, kecuali tangkapan layar produk.
"""
from __future__ import annotations

import datetime as dt
import os
from dataclasses import dataclass, field

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor

from . import iso as isomod
from . import measure, xmlfx
from .markup import parse
from .theme import Theme, for_mode

IN = 914400


def emu(x_in: float) -> int:
    return int(round(x_in * IN))


@dataclass
class Ctx:
    """Konteks build: token, fakta dari snapshot, folder md, dan peringatan yang dikumpulkan (teks hampir pasti meluap, dll)."""
    tokens: dict
    facts: dict
    base_dir: str
    font_profile: str | None = None
    warnings: list = field(default_factory=list)
    slide_no: int = 0
    n_slides: int = 12

    def warn(self, msg: str):
        self.warnings.append(f"slide {self.slide_no}: {msg}")

    def theme(self, mode: str) -> Theme:
        return for_mode(self.tokens, mode, self.font_profile)


# ───────────────────────────── presentasi dasar

def _rescale(shapes, f: float):
    for sh in shapes:
        sp = sh._element.find(qn("p:spPr"))
        if sp is not None and sp.find(qn("a:xfrm")) is not None:
            sh.left = int(sh.left * f)
            sh.width = int(sh.width * f)


def patch_theme(prs, tokens: dict, font_profile: str | None):
    """Tulis palet + font Fabius ke tema (supaya slide baru dan gaya turunan ikut)."""
    pal = tokens["palette"]
    prof = font_profile or tokens["fonts"].get("profile", "brand")
    fonts = tokens["fonts"][prof]
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    root.set("name", "Fabius")
    cs = root.find(f".//{A}clrScheme")
    cs.set("name", "Fabius")
    mapping = {"dk1": pal["night"], "lt1": pal["lav"], "dk2": pal["night2"], "lt2": pal["lav2"], "accent1": pal["violet"],
               "accent2": pal["violet2"], "accent3": pal["lav3"], "accent4": pal["mist"], "accent5": pal["gap"],
               "accent6": pal["indigo"], "hlink": pal["violet2"], "folHlink": pal["lav3"]}
    for key, hexv in mapping.items():
        node = cs.find(f"{A}{key}")
        for ch in list(node):
            node.remove(ch)
        etree.SubElement(node, f"{A}srgbClr").set("val", hexv)
    fs = root.find(f".//{A}fontScheme")
    fs.set("name", "Fabius")
    fs.find(f"{A}majorFont/{A}latin").set("typeface", fonts["display"])
    fs.find(f"{A}minorFont/{A}latin").set("typeface", fonts["body"])
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    # deck gelap: tx1 = terang, bg1 = gelap
    cm = prs.slide_master._element.find(qn("p:clrMap"))
    cm.set("bg1", "dk1"); cm.set("tx1", "lt1"); cm.set("bg2", "dk2"); cm.set("tx2", "lt2")


def build_time(when: dt.datetime | None = None) -> dt.datetime:
    """Cap waktu berkas: `when` bila diberi; kalau tidak `SOURCE_DATE_EPOCH` (konvensi build yang bisa diulang); kalau tidak, sekarang (UTC)."""
    if when is not None:
        return when
    epoch = os.environ.get("SOURCE_DATE_EPOCH", "")
    if epoch.isdigit():
        return dt.datetime.fromtimestamp(int(epoch), dt.timezone.utc).replace(tzinfo=None)
    return dt.datetime.now(dt.timezone.utc).replace(tzinfo=None, microsecond=0)


def base_presentation(tokens: dict, font_profile: str | None = None, when: dt.datetime | None = None) -> Presentation:
    prs = Presentation()
    old_w = prs.slide_width
    w_in, h_in = tokens["deck"]["size_in"]
    prs.slide_width, prs.slide_height = emu(w_in), emu(h_in)
    f = prs.slide_width / old_w
    _rescale(prs.slide_master.shapes, f)
    for lay in prs.slide_layouts:
        _rescale(lay.shapes, f)
    patch_theme(prs, tokens, font_profile)
    d = tokens["deck"]
    cp = prs.core_properties
    cp.title, cp.subject, cp.keywords = d["title"], d["subject"], d["keywords"]
    cp.author = cp.last_modified_by = d["author"]
    cp.comments = "Generated from vault/13-SubmmisionsDocs (one markdown file per slide) by build/build_deck.py"
    cp.language = d.get("language", "en-US")
    cp.created = cp.modified = build_time(when)
    return prs


def set_background(slide, theme: Theme):
    fill = slide.background.fill
    fill.solid()
    bgPr = slide._element.find(".//" + qn("p:bgPr"))
    for ch in list(bgPr):
        if ch.tag != qn("a:effectLst"):
            bgPr.remove(ch)
    g = etree.Element(qn("a:gradFill"), rotWithShape="1")
    lst = etree.SubElement(g, qn("a:gsLst"))
    for pos, hexv, al in theme.bg_stops:
        gs = etree.SubElement(lst, qn("a:gs"), pos=str(int(pos * 1000)))
        gs.append(xmlfx.color(hexv, al))
    etree.SubElement(g, qn("a:lin"), ang=str(90 * 60000), scaled="0")
    bgPr.insert(0, g)


def new_slide(prs, theme: Theme):
    s = prs.slides.add_slide(prs.slide_layouts[5])  # "Title Only": judul sungguhan untuk outline dan pembaca layar
    set_background(s, theme)
    return s


def set_transition(slide, kind: str | None):
    if not kind or kind == "none":
        return
    sld = slide._element
    tr = etree.Element(qn("p:transition"))
    tr.set("spd", "med")
    etree.SubElement(tr, qn("p:" + {"fade": "fade", "push": "push", "wipe": "wipe"}.get(kind, "fade")))
    sld.find(qn("p:clrMapOvr")).addnext(tr)


def set_notes(slide, text: str):
    slide.notes_slide.notes_text_frame.text = text


# ───────────────────────────── teks

@dataclass
class TS:
    """Gaya teks."""
    role: str = "body"            # display | body | mono
    size: float = 16
    color: str | None = None
    bold: bool = False
    italic: bool = False
    align: str = "l"
    line: float = 1.15
    after: float = 0
    caps: bool = False
    track: float = 0              # pt
    accent: str | None = None
    min_size: float | None = None


_ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
_ANCH = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def style(theme: Theme, kind: str, **over) -> TS:
    """Gaya baku per jenis teks (skala tipografi FE: display Archivo, body Inter, angka/hash JetBrains Mono)."""
    base = {
        "cover": TS("display", theme.size("cover"), line=0.92, bold=True),
        "title": TS("display", theme.size("title"), line=1.0),
        "headline": TS("display", theme.size("headline"), line=1.02),
        "sub": TS("body", theme.size("sub"), color=theme.muted, line=1.25),
        "lead": TS("body", theme.size("lead"), color=theme.muted, line=1.3),
        "body": TS("body", theme.size("body"), color=theme.text, line=1.3, after=6),
        "card_title": TS("display", theme.size("card_title"), bold=True, line=1.1, color=theme.strong),
        "card": TS("body", theme.size("card"), color=theme.muted, line=1.3),
        "caption": TS("body", theme.size("caption"), color=theme.muted, line=1.3),
        "tag": TS("mono", theme.size("tag"), color=theme.accent, caps=True, track=1.2, line=1.2),
        "mono": TS("mono", theme.size("mono"), color=theme.muted, line=1.25),
        "stat": TS("display", theme.size("stat"), bold=True, line=0.9, color=theme.strong),
        "stat_s": TS("display", theme.size("stat_s"), bold=True, line=0.95, color=theme.strong),
        "footer": TS("mono", theme.size("footer"), color=theme.faint, caps=True, track=1.0, line=1.2),
    }[kind]
    for k, v in over.items():
        setattr(base, k, v)
    if base.color is None:
        base.color = theme.text
    if base.accent is None:
        base.accent = theme.accent
    return base


def _rpr_extras(run, role: str, theme: Theme, track: float):
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", theme.tokens["deck"].get("language", "en-US"))
    if track:
        rPr.set("spc", str(int(round(track * 100))))
    latin = rPr.find(qn("a:latin"))
    if latin is not None:
        latin.set("pitchFamily", "49" if role == "mono" else "34")
        latin.set("charset", "0")


def _fill_tf(tf, text, theme: Theme, st: TS, anchor: str, size: float):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = _ANCH[anchor]
    paras = text if isinstance(text, (list, tuple)) else [text]
    for pi, ptxt in enumerate(paras):
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        p.alignment = _ALIGN[st.align]
        p.line_spacing = st.line
        p.space_before = Pt(0)
        p.space_after = Pt(st.after if pi < len(paras) - 1 else 0)
        src = str(ptxt).upper() if st.caps else str(ptxt)
        for li, runs in enumerate(parse(src)):
            if li > 0:
                br = p._p.add_br()
                brp = br.get_or_add_rPr()
                brp.set("sz", str(int(size * 100)))
            for r in runs:
                run = p.add_run()
                run.text = r.text
                role = "mono" if r.mono else st.role
                f = run.font
                f.name = theme.font(role)
                f.size = Pt(size)
                f.bold = bool(st.bold or r.bold or r.accent)
                f.italic = bool(st.italic or r.accent)
                col = st.accent if r.accent else (theme.accent if r.mono and st.role != "mono" else st.color)
                f.color.rgb = RGBColor.from_string(col)
                _rpr_extras(run, role, theme, st.track)


def _fit(ctx: Ctx | None, name: str, text, w: float, h: float, st: TS, fit: bool):
    size = st.size
    if not fit:
        return size
    paras = text if isinstance(text, (list, tuple)) else [text]
    bold = st.bold
    mono = st.role == "mono"

    def height_at(s):
        tot = 0.0
        for i, pt in enumerate(paras):
            src = str(pt).upper() if st.caps else str(pt)
            hh, _ = measure.block_height_in(src, w, s, st.line * (1.05 if st.role == "display" else 1.2) / 1.0, bold, mono, st.track)
            tot += hh + (st.after / 72.0 if i < len(paras) - 1 else 0)
        return tot

    floor = st.min_size or max(10.0, size * 0.6)
    s = size
    while s > floor and height_at(s) > h + 1e-6:
        s -= 0.5
    if height_at(s) > h + 1e-6 and ctx is not None:
        ctx.warn(f"teks mungkin meluap di '{name}' ({s:g} pt, kotak {w:.2f}x{h:.2f} in)")
    elif s < size * 0.88 and ctx is not None:
        ctx.warn(f"teks '{name}' menyusut dari {size:g} ke {s:g} pt agar muat (kotak {w:.2f}x{h:.2f} in)")
    return max(s, floor) if s < floor else s


def add_text(shapes, x, y, w, h, text, theme: Theme, st: TS, anchor: str = "t", name: str = "Text", ctx: Ctx | None = None,
             fit: bool = True, balance: bool = False):
    tb = shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    tb.name = name
    size = _fit(ctx, name, text, w, h, st, fit)
    if balance and isinstance(text, str):
        text = measure.balance(text, w, size, st.bold, st.role == "mono", st.track)
    _fill_tf(tb.text_frame, text, theme, st, anchor, size)
    return tb


def headline(slide, text, x, y, w, h, theme: Theme, st: TS, anchor: str = "t", ctx: Ctx | None = None, fit: bool = True):
    """Judul slide memakai placeholder judul (outline + aksesibilitas), gayanya dari `st`."""
    ph = slide.shapes.title
    ph.left, ph.top, ph.width, ph.height = emu(x), emu(y), emu(w), emu(h)
    ph.name = "Title"
    size = _fit(ctx, "Title", text, w, h, st, fit)
    tf = ph.text_frame
    tf.clear()
    _fill_tf(tf, text, theme, st, anchor, size)
    return ph


# ───────────────────────────── bentuk dasar

def _shape(shapes, kind, x, y, w, h, name):
    sh = shapes.add_shape(kind, emu(x), emu(y), emu(w), emu(h))
    sh.name = name
    sh.shadow.inherit = False
    return sh


def card(shapes, x, y, w, h, theme: Theme, radius: float = 0.2, name: str = "Card", fill=None, border=None, shadow=True,
         dash: str | None = None):
    """Panel kaca (Glass / Soft-depth): gradien putih transparan + tepi tipis + bayangan lembut di mode terang."""
    sh = _shape(shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, name)
    sh.adjustments[0] = min(0.5, radius / min(w, h))
    stops = fill if fill is not None else theme.card_stops
    xmlfx.gradient(sh, stops, angle=135)
    b = border if border is not None else theme.card_border
    xmlfx.line(sh, b[0], 0.75, b[1], dash=dash) if b else xmlfx.line(sh, None)
    xmlfx.effects(sh, shadow=theme.card_shadow if shadow else None)
    return sh


def rect(shapes, x, y, w, h, hexv, alpha=None, radius: float | None = None, name="Shape", line=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = _shape(shapes, kind, x, y, w, h, name)
    if radius:
        sh.adjustments[0] = min(0.5, radius / min(w, h))
    xmlfx.solid(sh, hexv, alpha)
    xmlfx.line(sh, *(line or (None,)))
    return sh


def dot(shapes, cx, cy, d, hexv, alpha=1.0, glow=None, name="Dot", line=None):
    sh = _shape(shapes, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, name)
    xmlfx.solid(sh, hexv, alpha)
    xmlfx.line(sh, *(line or (None,)))
    if glow:
        xmlfx.effects(sh, glow=glow)
    return sh


def oval(shapes, x, y, w, h, hexv=None, alpha=None, line=None, name="Oval"):
    sh = _shape(shapes, MSO_SHAPE.OVAL, x, y, w, h, name)
    if hexv:
        xmlfx.solid(sh, hexv, alpha)
    else:
        xmlfx.no_fill(sh)
    xmlfx.line(sh, *(line or (None,)))
    return sh


def polyline(shapes, pts_in, hexv, alpha=0.8, width=1.25, dash=None, tail=None, head=None, glow=None, name="Line", closed=False,
             fill=None):
    pts = [(emu(a), emu(b)) for a, b in pts_in]
    fb = shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
    fb.add_line_segments(pts[1:], close=closed)
    sh = fb.convert_to_shape()
    sh.name = name
    if fill:
        xmlfx.solid(sh, fill[0], fill[1])
    else:
        xmlfx.no_fill(sh)
    xmlfx.line(sh, hexv, width, alpha, dash=dash, tail=tail, head=head)
    if glow:
        xmlfx.effects(sh, glow=glow)
    return sh


def halo(shapes, cx, cy, rx, ry, hexv="6E4BFF", peak=0.34, steps=18, name="Halo"):
    """Cahaya lembut = elips konsentris beralfa kecil (tampil sama di PowerPoint dan LibreOffice, tanpa bergantung pada gradien radial)."""
    grp = shapes.add_group_shape()
    grp.name = name
    a = 1 - (1 - peak) ** (1.0 / steps)
    for i in range(steps):
        t = i / (steps - 1)
        s = 1.0 - 0.92 * t
        o = oval(grp.shapes, cx - rx * s, cy - ry * s, 2 * rx * s, 2 * ry * s, hexv, a, name=f"{name} {i + 1}")
    return grp


def hrule(shapes, x, y, w, theme: Theme, name="Rule"):
    c, a = theme.rule
    return polyline(shapes, [(x, y), (x + w, y)], c, a, 0.75, name=name)


# ───────────────────────────── komponen bahasa visual FE

def pill(shapes, x, y, text, theme: Theme, kind: str = "neutral", h: float = 0.3, size: float = 11, dot_on: bool = True,
         name: str | None = None, min_w: float = 0.0, ctx: Ctx | None = None, caps: bool = True) -> float:
    """Pil status mono kapital. kind: live | built | gated | neutral | gap. Mengembalikan lebar (inci) supaya deret chip bisa mengalir."""
    p = theme.p
    night = theme.night
    colors = {
        "live": (p["violet"], 0.32 if night else 0.14, p["violet2"] if night else p["violet"], 0.95, theme.strong if night else p["ink"]),
        "built": (p["indigo"], 0.55 if night else 0.14, p["lav3"] if night else p["indigo"], 0.55, theme.text),
        "gated": (None, 0, theme.faint, 0.75, theme.muted if night else p["dim"]),
        "neutral": (p["white"], 0.08 if night else 0.65, p["white"] if night else p["violet"], 0.22 if night else 0.30, theme.text),
        "gap": (p["gap"], 0.10, p["gap"], 0.85, p["gap"] if not night else "FF9AA7"),
    }[kind]
    fill_hex, fill_a, line_hex, line_a, txt = colors
    tw = measure.text_width_in(text.upper() if caps else text, size, False, True, 1.0 if caps else 0.0)
    pad = 0.15
    w = max(min_w, tw + 2 * pad + (0.2 if dot_on else 0))
    sh = _shape(shapes, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, name or f"Chip · {text}")
    sh.adjustments[0] = 0.5
    if fill_hex:
        xmlfx.solid(sh, fill_hex, fill_a)
    else:
        xmlfx.no_fill(sh)
    xmlfx.line(sh, line_hex, 0.9, line_a, dash="dash" if kind == "gated" else None)
    if dot_on:
        d = 0.095
        cx, cy = x + pad + d / 2, y + h / 2
        if kind == "live":
            dot(shapes, cx, cy, d, "F6F2FF", 1.0, glow=(p["violet2"], 6, 0.6), name="Chip dot")
        elif kind == "built":
            dot(shapes, cx, cy, d, "4A38C8", 1.0, line=(p["lav3"], 0.6, 0.8), name="Chip dot")
        elif kind == "gap":
            dot(shapes, cx, cy, d, p["gap"], 1.0, name="Chip dot")
        else:
            oval(shapes, cx - d / 2, cy - d / 2, d, d, None, None, line=(theme.faint, 0.9, 0.9), name="Chip dot")
    tx = x + pad + (0.2 if dot_on else 0)
    add_text(shapes, tx, y, w - pad - (tx - x), h, text, theme, TS("mono", size, color=txt, caps=caps, track=1.0 if caps else 0.0, line=1.0), anchor="m",
             name="Chip text", ctx=ctx, fit=False)
    return w


def marker(shapes, x, y, num: str, label: str, theme: Theme, mats: dict, ctx: Ctx | None = None, w: float = 6.0):
    """Penanda bernomor FE: blok kaca mini + `01 — LABEL` (mono kapital, violet)."""
    isomod.mini_cube(shapes, x + 0.11, y + 0.01, 0.22, mats, "violet" if not theme.night else "open", name="Marker cube")
    add_text(shapes, x + 0.34, y - 0.02, w, 0.28, f"{num} — {label}", theme, style(theme, "tag"), anchor="m", name="Marker", ctx=ctx, fit=False)


def footer(slide, theme: Theme, n_total: int, left_text: str, ctx: Ctx | None = None):
    w_in = slide.part.package.presentation_part.presentation.slide_width / IN
    h_in = slide.part.package.presentation_part.presentation.slide_height / IN
    y = h_in - 0.42
    add_text(slide.shapes, 0.65, y, 8.5, 0.22, left_text, theme, style(theme, "footer"), anchor="m", name="Footer", ctx=ctx, fit=False)
    tb = slide.shapes.add_textbox(emu(w_in - 0.65 - 1.4), emu(y), emu(1.4), emu(0.22))
    tb.name = "Slide number"
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    st = style(theme, "footer")
    fld = etree.SubElement(p._p, qn("a:fld"), id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}", type="slidenum")
    rPr = etree.SubElement(fld, qn("a:rPr"), lang="en-US", sz=str(int(st.size * 100)), spc="100")
    sf = etree.SubElement(rPr, qn("a:solidFill"))
    sf.append(xmlfx.color(st.color))
    etree.SubElement(rPr, qn("a:latin"), typeface=theme.font("mono"), pitchFamily="49", charset="0")
    t = etree.SubElement(fld, qn("a:t"))
    t.text = "‹#›"
    r = p.add_run()
    r.text = f" / {n_total:02d}"
    r.font.size = Pt(st.size)
    r.font.name = theme.font("mono")
    r.font.color.rgb = RGBColor.from_string(st.color)
    _rpr_extras(r, "mono", theme, st.track)


def picture(shapes, path: str, x, y, w, h, radius: float = 0.14, border=None, shadow=None, alt: str = "", name: str = "Screenshot",
            anchor: str = "top", crop=None):
    """Gambar dipotong menutupi kotak (tidak melar), sudut membulat, tepi tipis + bayangan opsional.

    `crop` = (kiri, atas, kanan, bawah) sebagai pecahan 0..1 yang dibuang lebih dulu (mis. membuang baris nav/chip dari tangkapan layar).
    """
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    pic = shapes.add_picture(path, emu(x), emu(y), emu(w), emu(h))
    pic.name = name
    l, t, r, b = (crop or (0, 0, 0, 0))
    vis_w, vis_h = 1 - l - r, 1 - t - b
    ar_v = (iw * vis_w) / (ih * vis_h)
    box_ar = w / h
    if ar_v > box_ar:   # area terlihat lebih lebar: potong kiri-kanan secara simetris
        extra = (1 - box_ar / ar_v) * vis_w
        l, r = l + extra / 2, r + extra / 2
    else:               # lebih tinggi: potong bawah (anchor top) atau simetris
        extra = (1 - ar_v / box_ar) * vis_h
        if anchor == "top":
            b += extra
        else:
            t, b = t + extra / 2, b + extra / 2
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = l, t, r, b
    pic.auto_shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
    prst = pic._element.spPr.find(qn("a:prstGeom"))
    av = prst.find(qn("a:avLst"))
    if av is None:
        av = etree.SubElement(prst, qn("a:avLst"))
    for ch in list(av):
        av.remove(ch)
    gd = etree.SubElement(av, qn("a:gd"))
    gd.set("name", "adj")
    gd.set("fmla", f"val {int(min(0.5, radius / min(w, h)) * 100000)}")
    if border:
        xmlfx.line(pic, border[0], border[1], border[2])
    if shadow:
        xmlfx.effects(pic, shadow=shadow)
    xmlfx.set_descr(pic, alt, name)
    return pic
